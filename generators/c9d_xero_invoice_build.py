#!/usr/bin/env python3
"""c9d_xero_invoice_build.py: the C9D Consulting invoice, built as a Xero custom .docx
branding theme (The Dossier, print mode). Spec: c9d-consulting-brand/templates/xero-invoice-spec.md.

Usage:
  python3 generators/c9d_xero_invoice_build.py template out.docx   # Xero merge fields, for upload
  python3 generators/c9d_xero_invoice_build.py sample   out.docx   # sample values only, for review

In the template, every merge field shows a sample value as its field result, so the unmerged
file reads like a finished invoice in Word and in previewers; Xero reads only the field codes.

Tokens, attribution and closing line come from the brand repo (tokens.json) through
c9d_docx_build.load_tokens, so nothing about the palette is compiled in here.

Xero renders custom .docx themes itself and strips embedded fonts, so this theme is set
in faces every renderer has: Arial stands in for Inter Tight and Courier New for
JetBrains Mono. Change FACE below if Xero is confirmed to honour the brand faces.
"""
import os, sys, copy
HERE = os.path.dirname(os.path.abspath(__file__))
BRAND_REPO = os.path.abspath(os.path.join(HERE, ".."))
sys.path.insert(0, HERE)

import c9d_docx_build as b
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

b.load_tokens(os.path.join(BRAND_REPO, "c9d-consulting-brand", "assets", "tokens.json"))
C, BRAND = b.C, b.BRAND

FACE = dict(sans="Arial", mono="Courier New")
LOGO_SRC = os.path.join(BRAND_REPO, "c9d-consulting-brand", "assets", "logos",
                        "c9d-consulting-wordmark-horizontal-light-transparent.png")
import tempfile
LOGO = os.path.join(tempfile.gettempdir(), "c9d-wordmark-light-cropped.png")


def prepare_logo():
    """Crop the brand PNG to its ink and flatten onto white: Xero rasterises images and
    handles transparency poorly, and the padding would inset the wordmark from the margin."""
    from PIL import Image
    im = Image.open(LOGO_SRC)
    c = im.crop(im.getbbox())
    bg = Image.new("RGB", c.size, (255, 255, 255)); bg.paste(c, mask=c.split()[3])
    bg.save(LOGO)
SIDE = 0.9
CONTENT = 8.5 - 2 * SIDE   # 6.7in text block
PAD = 144                  # 0.1in cell side padding, as the brand build

# ---- sample data for the review render ----------------------------------------------
SAMPLE = {
    "InvoiceNumber": "INV-0001",
    "InvoiceDate": "Nov 2, 2026",
    "DueDate": "Nov 17, 2026",
    "Reference": "ENG-2026-001 · SOW-RT-01 · PO 44817",
    "ContactName": "Northwind Holdings, Inc.",
    "PostalAddress": "Attn: Accounts Payable\n500 Main Street, Suite 1200\nDallas, TX 75201",
    "OrganisationName": "C9D Consulting LLC",
    "OrganisationPostalAddress": "Dallas, TX",
    "SubTotal": "38,850.00",
    "TaxTotal": "0.00",
    "Total": "38,850.00",
    "AmountPaid": "0.00",
    "AmountDue": "38,850.00",
    # results shown for the single line-item row in the unmerged template
    "Description": "Retainer delivery team, November 2026",
    "Quantity": "1.00",
    "UnitAmount": "38,850.00",
    "LineAmount": "38,850.00",
}
SAMPLE_LINES = [
    ("Retainer delivery team, November 2026\nSOW-RT-01, period of Nov 1 to Nov 30, in advance", "1.00", "37,000.00", "37,000.00"),
    ("Technology and administrative fee (5%)\nCSA 4.03", "1.00", "1,850.00", "1,850.00"),
    ("Third-party costs at cost, no markup\nCSA 4.04, receipts attached", "1.00", "0.00", "0.00"),
]

DATE = r'\@ "MMM d, yyyy"'
MONEY = r'\# "#,##0.00"'
QTY = r'\# "#,##0.00"'


class Inv:
    def __init__(self, mode):
        self.mode = mode
        self.doc = Document()
        self._setup()

    # ---- primitives ------------------------------------------------------------------
    def _setup(self):
        d = self.doc
        s = d.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(SIDE)
        s.top_margin, s.bottom_margin = Inches(0.6), Inches(0.95)
        s.header_distance, s.footer_distance = Inches(0.35), Inches(0.4)
        st = d.styles["Normal"]
        st.font.name = FACE["sans"]; st.font.size = Pt(9.5)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.line_spacing = 1.0
        # docDefaults in the stock template carry 10pt after and 1.15 lines; zero them
        dd = d.styles.element.find(qn("w:docDefaults"))
        if dd is not None:
            for sp in dd.iter(qn("w:spacing")):
                sp.set(qn("w:after"), "0"); sp.set(qn("w:line"), "240"); sp.set(qn("w:lineRule"), "auto")
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FACE["sans"])
        # no stray first paragraph
        body = d.element.body
        for p in list(body.findall(qn("w:p"))):
            body.remove(p)

    def run(self, p, text, size=9.5, color=None, mono=False, bold=False, caps=False, spacing=None):
        r = p.add_run(text)
        b.set_run_font(r, FACE["mono"] if mono else FACE["sans"], size, color or C["ink"],
                       bold=bold, caps=caps, spacing=spacing)
        return r

    def field(self, p, name, switch="", size=9.5, color=None, mono=False, bold=False):
        """A Xero merge field («Name») in template mode, the sample value in sample mode."""
        if self.mode == "sample":
            val = SAMPLE.get(name, name)
            lines = val.split("\n")
            for i, ln in enumerate(lines):
                r = self.run(p, ln, size, color, mono, bold)
                if i < len(lines) - 1:
                    r.add_break()
            return
        # The field result is what Word, previewers and PDF exports show until Xero merges.
        # Showing realistic sample values (not «FieldName») keeps the unmerged template
        # readable and laid out exactly like the invoice; Xero only reads the field code.
        self._merge(p, f"MERGEFIELD {name} {switch} \\* MERGEFORMAT".replace("  ", " "),
                    SAMPLE.get(name, name), size, color, mono, bold)

    def _merge(self, p, instr, result, size, color=None, mono=False, bold=False):
        r = self.run(p, "x", size, color, mono, bold)
        b.add_field(r, instr)
        # the add_field result run carries "1"; replace it with the result text (multi-line ok)
        res = None
        for t in p._p.iter(qn("w:t")):
            if t.text == "1":
                res = t; break
        run_el = res.getparent()
        lines = result.split("\n") if result else [""]
        res.text = lines[0]
        res.set(qn("xml:space"), "preserve")
        for ln in lines[1:]:
            br = OxmlElement("w:br"); run_el.append(br)
            t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = ln; run_el.append(t)
        if not result:
            run_el.remove(res)

    def raw_field(self, p, name, size=9.5):
        """TableStart / TableEnd markers: empty result, so they never show in Word or a preview."""
        if self.mode == "sample":
            return
        self._merge(p, f"MERGEFIELD {name} \\* MERGEFORMAT", "", size, C["ink_dim"])

    def para(self, container=None, before=0, after=0, line=1.25, align=None, keep_next=False):
        p = (container or self.doc).add_paragraph()
        b.para_format(p, before=before, after=after, line=line, keep_next=keep_next, align=align)
        return p

    def table(self, rows, widths, container=None):
        t = (container or self.doc).add_table(rows=rows, cols=len(widths))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        b.table_no_borders(t)
        b.table_fixed(t, [Inches(w) for w in widths])
        tw = t._tbl.tblPr.find(qn("w:tblW"))
        # percentage width: simple previewers (iOS Quick Look) autofit dxa tables to their
        # content; a 100% table holds the full text block there and in Word
        tw.set(qn("w:type"), "pct"); tw.set(qn("w:w"), str(int(round(sum(widths) / CONTENT * 5000))))
        return t

    @staticmethod
    def align_fill(t, pad=PAD):
        """Default cell margins plus tblInd equal to the left padding, so a filled or
        ruled table sits exactly on the text margins in Word and LibreOffice."""
        tblPr = t._tbl.tblPr
        ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), str(pad)); ind.set(qn("w:type"), "dxa")
        tblPr.append(ind)
        tcm = OxmlElement("w:tblCellMar")
        for side, v in (("top", 0), ("left", pad), ("bottom", 0), ("right", pad)):
            e = OxmlElement("w:" + side); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); tcm.append(e)
        tblPr.append(tcm)

    @staticmethod
    def cell_pad(cell, top=90, bottom=60, left=PAD, right=PAD):
        b.cell_margins(cell, top=top, bottom=bottom, left=left, right=right)

    @staticmethod
    def first_para(cell):
        return cell.paragraphs[0]

    def label(self, p, text, color=None):
        self.run(p, text, 7.5, color or C["stamp"], mono=True, caps=True, spacing=0.12)

    # ---- sections --------------------------------------------------------------------
    def chrome(self):
        t = self.table(1, [2.4, 2.3, 2.0])
        b.table_zero_margins(t)
        c0, c1, c2 = t.rows[0].cells
        p = self.first_para(c0); b.para_format(p, line=1.0)
        self.run(p, "FILE ", 7.5, C["ink_dim"], mono=True, spacing=0.08)
        self.field(p, "InvoiceNumber", size=7.5, color=C["ink"], mono=True)
        p = self.first_para(c1); b.para_format(p, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.run(p, "CLIENT CONFIDENTIAL", 7.5, C["stamp"], mono=True, spacing=0.12)
        p = self.first_para(c2); b.para_format(p, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        self.run(p, "§ INVOICE", 7.5, C["ink_dim"], mono=True, spacing=0.08)
        for c in (c0, c1, c2):
            b.cell_borders(c, bottom=(C["rule_bright"], 4))
            b.cell_margins(c, top=0, bottom=70, left=0, right=0)

    def masthead(self):
        t = self.table(1, [3.5, 3.2])
        b.table_zero_margins(t)
        c0, c1 = t.rows[0].cells
        for c in (c0, c1):
            b.cell_margins(c, top=0, bottom=0, left=0, right=0)
        p = self.first_para(c0); b.para_format(p, before=16, after=2, line=1.0)
        p.add_run().add_picture(LOGO, width=Inches(2.1))
        for inl in p._p.iter(qn("wp:inline")):
            for a in ("distT", "distB", "distL", "distR"):
                inl.set(a, "0")
        b.cell_valign(c0, "bottom")
        p = self.first_para(c1); b.para_format(p, before=12, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        self.run(p, "Invoice", 28, C["ink_bright"], spacing=-0.03)
        b.cell_valign(c1, "bottom")

    def parties(self):
        p = self.para(before=0, after=0, line=1.0)
        b.add_bottom_border(p, C["rule"], sz=4, space=1)
        p.paragraph_format.space_after = Pt(12)
        widths = [2.45, 2.15, 2.1]
        t = self.table(1, widths)
        b.table_zero_margins(t)
        c0, c1, c2 = t.rows[0].cells
        # Billed to
        p = self.first_para(c0); b.para_format(p, after=3, line=1.0); self.label(p, "Billed to")
        p = c0.add_paragraph(); b.para_format(p, after=1, line=1.3)
        self.field(p, "ContactName", size=10, color=C["ink_bright"], bold=True)
        p = c0.add_paragraph(); b.para_format(p, line=1.3)
        self.field(p, "PostalAddress", size=9, color=C["ink_mute"])
        # From
        p = self.first_para(c1); b.para_format(p, after=3, line=1.0); self.label(p, "From")
        p = c1.add_paragraph(); b.para_format(p, after=1, line=1.3)
        self.field(p, "OrganisationName", size=10, color=C["ink_bright"], bold=True)
        p = c1.add_paragraph(); b.para_format(p, line=1.3)
        self.field(p, "OrganisationPostalAddress", size=9, color=C["ink_mute"])
        p = c1.add_paragraph(); b.para_format(p, line=1.3)
        self.run(p, "c9d.consulting", 9, C["stamp"])
        # Invoice facts
        rows = [("Invoice no.", "InvoiceNumber", "", True),
                ("Issued", "InvoiceDate", DATE, False),
                ("Due", "DueDate", DATE, False),
                ("Reference", "Reference", "", False)]
        for i, (lab, name, sw, mono) in enumerate(rows):
            p = self.first_para(c2) if i == 0 else c2.add_paragraph()
            b.para_format(p, after=0 if i < 3 else 0, line=1.0)
            self.label(p, lab, C["ink_dim"])
            p = c2.add_paragraph(); b.para_format(p, after=4, line=1.15)
            self.field(p, name, sw, size=9.5 if not mono else 9.5,
                       color=C["ink_bright"], mono=mono, bold=(name == "DueDate"))
        for c in (c0, c1, c2):
            b.cell_valign(c, "top")
        b.cell_margins(c0, top=0, bottom=0, left=0, right=PAD)
        b.cell_margins(c1, top=0, bottom=0, left=PAD, right=PAD)
        b.cell_margins(c2, top=0, bottom=0, left=PAD, right=0)

    def amount_band(self):
        p = self.para(after=8, line=1.0)
        t = self.table(1, [3.35, 3.35])
        self.align_fill(t, 200)
        c0, c1 = t.rows[0].cells
        for c in (c0, c1):
            b.cell_shade(c, C["bg2"])
            b.cell_borders(c, top=(C["stamp"], 12))
        self.cell_pad(c0, top=120, bottom=120, left=200, right=PAD)
        self.cell_pad(c1, top=120, bottom=120, left=PAD, right=200)
        b.cell_valign(c0, "center"); b.cell_valign(c1, "center")
        p = self.first_para(c0); b.para_format(p, after=4, line=1.0)
        self.label(p, "Amount due · USD")
        p = c0.add_paragraph(); b.para_format(p, line=1.0)
        self.run(p, "$", 24, C["stamp"], spacing=-0.02)
        self.field(p, "AmountDue", MONEY, size=24, color=C["ink_bright"])
        p = self.first_para(c1); b.para_format(p, after=4, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        self.label(p, "Payment due", C["ink_dim"])
        p = c1.add_paragraph(); b.para_format(p, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        self.field(p, "DueDate", DATE, size=14, color=C["ink_bright"])
        p = c1.add_paragraph(); b.para_format(p, before=4, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        self.run(p, "ACH preferred · quote ", 8, C["ink_mute"])
        self.field(p, "InvoiceNumber", size=8, color=C["ink_mute"], mono=True)
        b.row_no_split(t.rows[0])

    def line_items(self):
        p = self.para(before=0, after=8, line=1.0)
        widths = [3.85, 0.85, 1.0, 1.0]
        n = 1 if self.mode == "template" else len(SAMPLE_LINES)
        t = self.table(1 + n, widths)
        self.align_fill(t)
        heads = ["Description", "Qty / hrs", "Rate", "Amount"]
        for i, h in enumerate(heads):
            c = t.rows[0].cells[i]
            b.cell_shade(c, C["bg2"])
            b.cell_borders(c, bottom=(C["rule_bright"], 4))
            # paint the vertical seams in the fill colour so no hairline gap shows between cells
            for side in ("left", "right"):
                e = c._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:" + side))
                e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "2"); e.set(qn("w:color"), C["bg2"]); e.set(qn("w:space"), "0")
            self.cell_pad(c, top=80, bottom=60)
            b.cell_valign(c, "center")
            p = self.first_para(c)
            b.para_format(p, line=1.0, align=None if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
            self.run(p, h.upper(), 7.5, C["ink_mute"], mono=True, spacing=0.1)
        b.row_no_split(t.rows[0], header=True)
        rows = SAMPLE_LINES if self.mode == "sample" else [None]
        for ri, data in enumerate(rows, start=1):
            cells = t.rows[ri].cells
            for i, c in enumerate(cells):
                b.cell_borders(c, bottom=(C["rule"], 4))
                self.cell_pad(c, top=100, bottom=80)
                b.cell_valign(c, "top")
                p = self.first_para(c)
                b.para_format(p, line=1.3, align=None if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
                if data is not None:
                    if i == 0:
                        first, *rest = data[0].split("\n")
                        self.run(p, first, 9.5, C["ink_bright"])
                        for ln in rest:
                            r = self.run(p, "", 9.5); r.add_break()
                            self.run(p, ln, 8.5, C["ink_mute"])
                    else:
                        self.run(p, data[i], 9.5, C["ink"] if i < 3 else C["ink_bright"],
                                 mono=False)
                else:
                    if i == 0:
                        self.raw_field(p, "TableStart:LineItem", 9.5)
                        self.field(p, "Description", size=9.5, color=C["ink_bright"])
                    elif i == 1:
                        self.field(p, "Quantity", QTY, size=9.5)
                    elif i == 2:
                        self.field(p, "UnitAmount", MONEY, size=9.5)
                    else:
                        self.field(p, "LineAmount", MONEY, size=9.5, color=C["ink_bright"])
                        self.raw_field(p, "TableEnd:LineItem", 9.5)
            b.row_no_split(t.rows[ri])

    def totals(self):
        """Settlement block: remittance details on the left, totals on the right, one table
        so the two stay together and the page holds everything on one sheet."""
        widths = [1.15, 2.0, 0.55, 1.9, 1.1]
        bank = [(None, "How to pay \u00b7 ACH or wire"),
                ("Beneficiary", "C9D Consulting LLC"),
                ("Bank", "[Bank name]"),
                ("Routing (ACH)", "[ACH routing number]"),
                ("Routing (wire)", "[Wire routing number]"),
                ("Account", "[Account number]"),
                ("Reference", "InvoiceNumber")]
        tots = [("Subtotal", "SubTotal", False),
                ("Sales tax", "TaxTotal", False),
                ("Total", "Total", False),
                ("Less amount paid", "AmountPaid", False),
                ("Amount due (USD)", "AmountDue", True)]
        self.para(after=10, line=1.0)
        t = self.table(len(bank), widths)
        self.align_fill(t)
        for ri in range(len(bank)):
            r = t.rows[ri]
            c = r.cells
            off = len(bank) - len(tots)
            ti = ri - off
            strong = 0 <= ti and tots[ti][2]
            for ci, cell in enumerate(c):
                top, bot = (110, 80) if strong else (60, 45)
                left = 0 if ci == 0 else PAD
                if ci == 1:
                    left = 60
                self.cell_pad(cell, top=top, bottom=bot, left=left, right=PAD if ci != 1 else 0)
                b.cell_valign(cell, "bottom" if ci >= 3 else "top")
            # left: remittance
            k, v = bank[ri]
            if k is None:
                p = self.first_para(c[0]); b.para_format(p, line=1.0)
                self.label(p, v)
                c[0].merge(c[1])
            else:
                for cell in (c[0], c[1]):
                    b.cell_borders(cell, bottom=(C["rule"], 2))
                p = self.first_para(c[0]); b.para_format(p, line=1.0)
                self.run(p, k, 8, C["ink_mute"])
                p = self.first_para(c[1]); b.para_format(p, line=1.0)
                if v == "InvoiceNumber":
                    self.field(p, v, size=8.5, color=C["stamp"], mono=True, bold=True)
                else:
                    ph = v.startswith("[")
                    self.run(p, v, 8.5, C["stamp"] if ph else C["ink_bright"])
            # right: totals
            if ti >= 0:
                lab, name, strong = tots[ti]
                if strong:
                    for cell in (c[3], c[4]):
                        b.cell_borders(cell, top=(C["ink_bright"], 8))
                elif name == "Total":
                    for cell in (c[3], c[4]):
                        b.cell_borders(cell, top=(C["rule"], 4))
                p = self.first_para(c[3]); b.para_format(p, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
                if strong:
                    self.run(p, lab.upper(), 8, C["stamp"], mono=True, spacing=0.1)
                else:
                    self.run(p, lab, 9, C["ink_mute"])
                p = self.first_para(c[4]); b.para_format(p, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
                self.field(p, name, MONEY, size=12 if strong else 9.5, color=C["ink_bright"], bold=strong)
            b.row_no_split(r)
            for cell in r.cells:
                for pp in cell.paragraphs:
                    pp.paragraph_format.keep_with_next = ri < len(bank) - 1

    def terms(self):
        p = self.para(before=14, after=4, line=1.0, keep_next=True)
        self.label(p, "Terms")
        p = self.para(after=0, line=1.4)
        self.run(p, "Payment is due by the date above. Quote the invoice number as the payment reference so it "
                    "reconciles on arrival. Invoiced under the Consulting Services Agreement and the Statement of "
                    "Work named in the reference line; balances unpaid after the due date carry the late charge set "
                    "out in that agreement. Billing questions: ", 8.5, C["ink"])
        self.run(p, "[billing email]", 8.5, C["stamp"])
        self.run(p, ".", 8.5, C["ink"])

    def closing(self):
        p = self.para(before=14, after=0, line=1.0)
        b.add_top_border(p, C["rule"], sz=4, space=6)
        self.run(p, "→ " + BRAND["closing"].rstrip(".").upper(), 8, C["stamp"], mono=True, bold=True, spacing=0.18)
        p.paragraph_format.keep_with_next = False

    def footer(self):
        f = self.doc.sections[0].footer
        f.is_linked_to_previous = False
        for p in list(f.paragraphs):
            p._p.getparent().remove(p._p)
        t = f.add_table(rows=1, cols=3, width=Inches(CONTENT))
        b.table_no_borders(t); b.table_zero_margins(t)
        b.table_fixed(t, [Inches(1.9), Inches(3.0), Inches(1.8)])
        tw = t._tbl.tblPr.find(qn("w:tblW"))
        tw.set(qn("w:type"), "pct"); tw.set(qn("w:w"), "5000")
        c0, c1, c2 = t.rows[0].cells
        for c in (c0, c1, c2):
            b.cell_borders(c, top=(C["rule"], 4))
            b.cell_margins(c, top=70, bottom=0, left=0, right=0)
        p = c0.paragraphs[0]; b.para_format(p, line=1.0)
        self.run(p, "C9D", 8, C["stamp"])
        self.run(p, " Consulting", 8, C["ink_bright"])
        p = c1.paragraphs[0]; b.para_format(p, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.run(p, BRAND["attribution"] + ".", 8, C["ink_mute"])
        p = c2.paragraphs[0]; b.para_format(p, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        self.run(p, "c9d.consulting", 7.5, C["ink_dim"], mono=True, spacing=0.04)
        tail = f.add_paragraph(); b.para_format(tail, line=0.5)

    def build(self, out):
        self.chrome()
        self.masthead()
        self.parties()
        self.amount_band()
        self.line_items()
        self.totals()
        self.terms()
        self.closing()
        self.footer()
        PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
               "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
               "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
               "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
               "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
               "sectPr", "pPrChange"]
        ftr = self.doc.sections[0].footer._element
        for root in (self.doc.element, ftr):
            b.schema_order(root)
            for el in root.iter(qn("w:pPr")):
                rank = {qn("w:" + n): i for i, n in enumerate(PPR)}
                kids = sorted(list(el), key=lambda k: rank.get(k.tag, len(PPR)))
                for k in kids: el.remove(k)
                for k in kids: el.append(k)
        z = self.doc.settings.element.find(qn("w:zoom"))
        if z is not None:
            z.set(qn("w:percent"), "100")
        self.doc.core_properties.title = "C9D Consulting Invoice"
        self.doc.core_properties.author = "C9D Consulting"
        self.doc.save(out)


if __name__ == "__main__":
    mode, out = sys.argv[1], sys.argv[2]
    assert mode in ("template", "sample")
    prepare_logo()
    Inv(mode).build(out)
    print("wrote", out)
