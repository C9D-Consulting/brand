#!/usr/bin/env python3
"""c9d_docx_build.py: build a C9D Consulting branded .docx (The Dossier, light mode)
from a markdown source with YAML front matter. Same source dialect and front matter as
c9d_pdf_build.py and c9d_html_build.py, so one source produces every format.

Usage:  python3 c9d_docx_build.py source.md out.docx [--no-embed] [--tokens PATH] [--check] [--draft]
Needs:  pip install python-docx fonttools
Tokens: the light palette, the attribution line and the closing line are read from the brand
        repo's tokens.json (--tokens PATH, $C9D_TOKENS, the repo copy, ./tokens.json beside this
        script, or fetched from github.com/C9D-Consulting/brand). Nothing is compiled in.
Fonts:  fetched from the Google Fonts repository on first run into ./fonts beside this
        script (Inter Tight, JetBrains Mono, Instrument Serif; all SIL OFL), instanced
        to static weights, subset, obfuscated and embedded per ECMA-376. Same cache and file
        names as c9d_pdf_build.py, so the two scripts share one ./fonts folder.
Check:  --check verifies the built file: word parity with the source, embedded brand faces,
        the attribution exactly as tokens.json holds it, and the voice rules. --draft allows
        unresolved [placeholders].

Layout v1.4: US Letter, 1.3in top and 1.0in side and bottom margins, 11pt body across the
full 6.5in text block, balanced table cells, a one-page cover, and the locked closing line
as the closing signature. The decisions are recorded in templates/docx-spec.md.

Front matter keys: file_id, classification, title, subtitle, prepared_for, prepared_by,
version, section_marker. Body: # title, ## "EYEBROW · Title" sections (a bare number
renders as SECTION 01), ### and #### subheads, GFM tables, - and 1. lists, **bold**,
*italic*, `mono`, [links](https://...), and [bracketed placeholders], which render in
stamp amber.
"""
import re, sys, os, io, uuid, zipfile, shutil
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from fontTools.ttLib import TTFont
from fontTools import subset as ftsubset

VERSION = "1.4"

# ---- tokens: read from the brand repo's tokens.json, never compiled in -------------------
# Same resolution order as c9d_pdf_build.py: --tokens PATH, $C9D_TOKENS, the repo copy when
# this script runs from the brand repo's generators/ folder, tokens.json beside this script,
# then the source of record on GitHub (cached beside this script). No fallback palette: if
# tokens.json cannot be found the build stops, so a token change is made once, in the repo.
TOKENS_URL = ("https://raw.githubusercontent.com/C9D-Consulting/brand/main/"
              "c9d-consulting-brand/assets/tokens.json")
HERE = os.path.dirname(os.path.abspath(__file__))
C = {}       # light palette, hex without '#', filled by load_tokens()
BRAND = {}   # attribution, closing line, typeface names, token source

def load_tokens(path=None, refresh=False):
    import json, urllib.request
    repo = os.path.join(HERE, "..", "c9d-consulting-brand", "assets", "tokens.json")
    cands = [path, os.environ.get("C9D_TOKENS"), repo, os.path.join(HERE, "tokens.json")]
    src = next((c for c in cands if c and os.path.exists(c)), None)
    cache = os.path.join(HERE, "tokens.json")
    if src is None or refresh:
        try:
            data = urllib.request.urlopen(TOKENS_URL, timeout=20).read()
            open(cache, "wb").write(data); src = cache
        except Exception as e:
            if src is None:
                sys.exit(f"tokens.json not found and could not be fetched ({e}). "
                         "Pass --tokens path/to/c9d-consulting-brand/assets/tokens.json.")
    d = json.load(open(src, encoding="utf8"))
    c = d["color"]["light"]   # Word documents are print artifacts: light mode
    v = lambda grp, k: c[grp][k]["value"].lstrip("#").lower()
    C.clear(); C.update(
        ink_bright=v("ink", "ink-bright"), ink=v("ink", "ink"), ink_mute=v("ink", "ink-mute"),
        ink_dim=v("ink", "ink-dim"), stamp=v("signal", "stamp"), stamp_dim=v("signal", "stamp-dim"),
        rule=v("rule", "rule"), rule_bright=v("rule", "rule-bright"),
        bg2=v("ground", "bg-2"), bg3=v("ground", "bg-3"), classified=v("crit", "classified"))
    st = d["typography"]["stack"]
    BRAND.update(
        source=src, token_version=d["$metadata"].get("version", "?"),
        attribution=d["naming"]["founder"]["attribution"],
        closing=d["naming"]["tagline"]["value"],
        families=(st["display"]["family"], st["mono"]["family"], st["italic-accent"]["family"]))
    known = {"Inter Tight", "JetBrains Mono", "Instrument Serif"}
    missing = [f for f in BRAND["families"] if f not in known]
    if missing:
        sys.exit(f"tokens.json names typefaces this script cannot embed: {missing}")
    return BRAND

F = dict(light="Inter Tight Light", regular="Inter Tight", medium="Inter Tight Medium",
         semibold="Inter Tight SemiBold", mono="JetBrains Mono Medium",
         mono_bold="JetBrains Mono SemiBold", serif_i="Instrument Serif")
FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
FONT_FILES = {
    F["light"]: "InterTightLight-Regular.ttf", F["regular"]: "InterTight-Regular.ttf",
    F["medium"]: "InterTightMedium-Regular.ttf", F["semibold"]: "InterTightSemiBold-Regular.ttf",
    F["mono"]: "JetBrainsMonoMedium-Regular.ttf", F["mono_bold"]: "JetBrainsMonoSemiBold-Regular.ttf",
    F["serif_i"]: "InstrumentSerif-Italic.ttf",
}
BODY_MEASURE_INDENT = Inches(0)     # v1.2: body uses the full text block, no right-hand measure indent
SIDE_MARGIN_IN = 1.0                 # standard 1in left and right margins
CONTENT_IN = 8.5 - 2 * SIDE_MARGIN_IN  # 6.5in text block
BODY_PT = 11
TOP_MARGIN_IN = 1.3                  # v1.3: clear air between the header hairline and the first line of body
COVER_GAP_PT = 48                    # v1.4: air above the cover title and above the cover metadata; sized so the cover fits one page even under font substitution
CELL_LINE = 1.15                     # v1.4: cell line height; 1.4 left dead leading under the text
CELL_PAD_TOP = 100                   # v1.4: cell top padding, twips (5pt)
CELL_PAD_BOTTOM = 40                 # v1.4: cell bottom padding, twips (2pt); the line box already carries the descent
RULE_GAP_PT = 6                      # v1.3: height of the paragraph carrying the heading rule
CELL_PAD_H = 144                     # v1.3: table cell left/right padding, twips (0.1in), every column


# ---- font bootstrap -----------------------------------------------------------
GF = "https://raw.githubusercontent.com/google/fonts/main/ofl"
GF_FILES = {
    "InterTight[wght].ttf": "intertight/InterTight%5Bwght%5D.ttf",
    "JetBrainsMono[wght].ttf": "jetbrainsmono/JetBrainsMono%5Bwght%5D.ttf",
    "InstrumentSerif-Italic.ttf": "instrumentserif/InstrumentSerif-Italic.ttf",
}

def ensure_fonts():
    """Download the three OFL families from the Google Fonts repo and instance the
    static weights The Dossier needs, as separate named families Word can use."""
    if all(os.path.exists(os.path.join(FONT_DIR, f)) for f in FONT_FILES.values()):
        return
    import urllib.request
    from fontTools.varLib import instancer
    src = os.path.join(FONT_DIR, "_src"); os.makedirs(src, exist_ok=True)
    for local, remote in GF_FILES.items():
        p = os.path.join(src, local)
        if not os.path.exists(p):
            urllib.request.urlretrieve(f"{GF}/{remote}", p)
    def rename(font, family, sub):
        full = family if sub == "Regular" else f"{family} {sub}"
        for rec in font["name"].names:
            nid = rec.nameID
            if nid == 1: rec.string = family
            elif nid == 2: rec.string = sub
            elif nid == 4: rec.string = full
            elif nid == 6: rec.string = (family + "-" + sub).replace(" ", "")
        font["name"].names = [r for r in font["name"].names if r.nameID not in (16, 17)]
    def inst(local, wght, family):
        f = TTFont(os.path.join(src, local))
        f = instancer.instantiateVariableFont(f, {"wght": wght})
        rename(f, family, "Regular")
        f["OS/2"].usWeightClass = wght
        f["OS/2"].fsSelection = (f["OS/2"].fsSelection & ~0x21) | 0x40
        f["head"].macStyle = 0
        f.save(os.path.join(FONT_DIR, FONT_FILES[family]))
    inst("InterTight[wght].ttf", 300, F["light"])
    inst("InterTight[wght].ttf", 400, F["regular"])
    inst("InterTight[wght].ttf", 500, F["medium"])
    inst("InterTight[wght].ttf", 600, F["semibold"])
    inst("JetBrainsMono[wght].ttf", 500, F["mono"])
    inst("JetBrainsMono[wght].ttf", 600, F["mono_bold"])
    shutil.copy(os.path.join(src, "InstrumentSerif-Italic.ttf"), os.path.join(FONT_DIR, FONT_FILES[F["serif_i"]]))

# ---- markdown parsing --------------------------------------------------------
def parse_front_matter(text):
    meta = {}
    if text.startswith("---"):
        end = text.index("\n---", 3)
        for line in text[3:end].strip().splitlines():
            if ":" in line:
                k, v = line.split(":", 1); meta[k.strip()] = v.strip()
        text = text[end + 4:]
    return meta, text

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\)|\[[^\]]+\])")
LINK_RE = re.compile(r"^\[([^\]]+)\]\(([^)\s]+)\)$")

def inline_runs(text):
    """Yield (kind, text) where kind in body|bold|italic|mono|placeholder."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            yield "bold", part[2:-2]
        elif part.startswith("*"):
            yield "italic", part[1:-1]
        elif part.startswith("`"):
            yield "mono", part[1:-1]
        elif LINK_RE.match(part):
            yield "link", part
        elif part.startswith("[") and part.endswith("]") and not part.startswith("[["):
            yield "placeholder", part
        else:
            yield "body", part

def parse_blocks(text):
    lines = text.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows)); continue
        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            blocks.append(("h%d" % len(m.group(1)), m.group(2).strip())); i += 1; continue
        if re.match(r"^- ", ln):
            items = []
            while i < len(lines) and re.match(r"^- ", lines[i]):
                items.append(lines[i][2:].strip()); i += 1
            blocks.append(("ul", items)); continue
        if re.match(r"^\d+\. ", ln):
            items = []
            while i < len(lines) and re.match(r"^\d+\. ", lines[i]):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i]).strip()); i += 1
            blocks.append(("ol", items)); continue
        para = [ln.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||- |\d+\. )", lines[i]):
            para.append(lines[i].strip()); i += 1
        blocks.append(("p", " ".join(para)))
    return blocks

# ---- low-level docx helpers --------------------------------------------------
def set_run_font(run, family, size, color, bold=False, italic=False, caps=False, spacing=None):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.insert(0, rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(a), family)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    run.font.italic = italic
    if caps:
        run.font.all_caps = True
    if spacing is not None:   # em fraction -> twentieths of a point
        sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), str(int(round(spacing * size * 20))))
        rPr.append(sp)

def para_format(p, before=0, after=0, line=1.6, keep_next=False, keep_together=True,
                right_indent=None, left_indent=None, hanging=None, align=None, page_break_before=False):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    pf.keep_together = keep_together
    pf.widow_control = True
    if page_break_before:
        pf.page_break_before = True
    if right_indent is not None: pf.right_indent = right_indent
    if left_indent is not None: pf.left_indent = left_indent
    if hanging is not None: pf.first_line_indent = -hanging
    if align is not None: pf.alignment = align

def add_hyperlink(p, text, url, size):
    """External link in stamp amber with no underline (docx-spec: never default blue)."""
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); h.append(r); p._p.append(h)
    from docx.text.run import Run
    run = Run(r, p); run.text = text
    set_run_font(run, F["regular"], size, C["stamp"])
    run.font.underline = False
    return run

def add_bottom_border(p, color, sz=4, space=1):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), str(space)); b.set(qn("w:color"), color)
    pbdr.append(b); pPr.append(pbdr)

def add_top_border(p, color, sz=4, space=1):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:top")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), str(space)); b.set(qn("w:color"), color)
    pbdr.append(b); pPr.append(pbdr)

def add_field(run, instr):
    """Complex field split across five runs, each carrying the source run's formatting,
    so Word and LibreOffice render the field result at the chrome size (v1.3)."""
    import copy
    r = run._r
    rpr = r.find(qn("w:rPr"))
    parts = []
    for kind in ("begin", "instr", "separate", "result", "end"):
        nr = OxmlElement("w:r")
        if rpr is not None:
            nr.append(copy.deepcopy(rpr))
        if kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; nr.append(it)
        elif kind == "result":
            t = OxmlElement("w:t"); t.text = "1"; nr.append(t)
        else:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); nr.append(fc)
        parts.append(nr)
    anchor = r
    for nr in parts:
        anchor.addnext(nr); anchor = nr
    r.getparent().remove(r)

def cell_margins(cell, top=120, bottom=120, left=0, right=280):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        e = OxmlElement("w:" + k); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcPr.append(mar)

TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
              "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]

def schema_order(root):
    """Reorder tcPr and tblPr children to the ECMA-376 sequence Word requires (v1.4)."""
    for tag, order in (("w:tcPr", TCPR_ORDER), ("w:tblPr", TBLPR_ORDER)):
        for el in root.iter(qn(tag)):
            kids = list(el)
            rank = {qn("w:" + n): i for i, n in enumerate(order)}
            kids.sort(key=lambda k: rank.get(k.tag, len(order)))
            for k in kids: el.remove(k)
            for k in kids: el.append(k)

def cell_valign(cell, where):
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign"); v.set(qn("w:val"), where); tcPr.append(v)

def cell_shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def cell_borders(cell, bottom=None, top=None):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + side)
        spec = {"bottom": bottom, "top": top}.get(side)
        if spec:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(spec[1])); e.set(qn("w:color"), spec[0]); e.set(qn("w:space"), "0")
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)

def table_no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + side); e.set(qn("w:val"), "nil"); b.append(e)
    tblPr.append(b)

def table_zero_margins(table):
    """Table-level cell margins of zero so chrome tables sit exactly on the text margins (v1.3)."""
    tcm = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + side); e.set(qn("w:w"), "0"); e.set(qn("w:type"), "dxa"); tcm.append(e)
    table._tbl.tblPr.append(tcm)

def table_fixed(table, widths):
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w
    grid = table._tbl.find(qn("w:tblGrid"))
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i].inches * 1440)))

def row_no_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit"); trPr.append(cs)
    if header:
        th = OxmlElement("w:tblHeader"); trPr.append(th)

# ---- writer ------------------------------------------------------------------
class Writer:
    def __init__(self, meta):
        self.meta = meta
        self.doc = Document()
        self.chars = set()
        st = self.doc.styles["Normal"]
        st.font.name = F["regular"]; st.font.size = Pt(BODY_PT)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), F["regular"])
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.top_margin = Inches(TOP_MARGIN_IN); sec.bottom_margin = Inches(1.0)
        sec.left_margin = sec.right_margin = Inches(SIDE_MARGIN_IN)
        sec.header_distance = sec.footer_distance = Inches(0.5)
        sec.different_first_page_header_footer = True
        self.content_width = Inches(CONTENT_IN)
        self.section_marker = meta.get("section_marker", "01")

    # -- runs --
    def run(self, p, text, kind="body", size=BODY_PT):
        if kind == "link":
            label, url = LINK_RE.match(text).groups()
            self.chars.update(label)
            return add_hyperlink(p, label, url, size)
        self.chars.update(text)
        r = p.add_run(text)
        if kind == "body":
            set_run_font(r, F["regular"], size, C["ink"])
        elif kind == "bold":
            set_run_font(r, F["semibold"], size, C["ink_bright"])
        elif kind == "italic":
            set_run_font(r, F["serif_i"], size + 0.5, C["ink"], italic=True)
        elif kind == "mono":
            set_run_font(r, F["mono"], size - 1, C["ink"])
        elif kind == "placeholder":
            set_run_font(r, F["regular"], size, C["stamp"])
        return r

    def runs(self, p, text, size=BODY_PT):
        for kind, t in inline_runs(text):
            self.run(p, t, kind, size)

    # -- chrome --
    def build_header_footer(self):
        sec = self.doc.sections[0]
        hdr = sec.header; hdr.is_linked_to_previous = False
        hp = hdr.paragraphs[0]; hp.text = ""
        t = hdr.add_table(rows=1, cols=3, width=self.content_width)
        table_no_borders(t); table_zero_margins(t); table_fixed(t, [Inches(2.0), Inches(CONTENT_IN - 4.0), Inches(2.0)])
        cells = t.rows[0].cells
        for c in cells:
            cell_margins(c, top=0, bottom=60, left=0, right=0)
        p = cells[0].paragraphs[0]; para_format(p, line=1.2)
        r = p.add_run(self.meta.get("file_id", "")); set_run_font(r, F["mono"], 8, C["ink_mute"], spacing=0.04)
        p = cells[1].paragraphs[0]; para_format(p, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(self.meta.get("classification", "")); set_run_font(r, F["medium"], 8, C["stamp"], caps=True, spacing=0.18)
        p = cells[2].paragraphs[0]; para_format(p, line=1.2, align=WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(f"§ {self.section_marker} · "); set_run_font(r, F["mono"], 8, C["ink_mute"], spacing=0.04)
        r = p.add_run(); set_run_font(r, F["mono"], 8, C["ink_mute"], spacing=0.04); add_field(r, "PAGE")
        r = p.add_run(" / "); set_run_font(r, F["mono"], 8, C["ink_mute"], spacing=0.04)
        r = p.add_run(); set_run_font(r, F["mono"], 8, C["ink_mute"], spacing=0.04); add_field(r, "NUMPAGES")
        # hairline under the header row
        hp2 = hdr.add_paragraph(); para_format(hp2, line=1.0); add_bottom_border(hp2, C["rule"], sz=4, space=0)
        set_run_font(hp2.add_run(""), F["regular"], 2, C["ink"])
        # remove the empty leading paragraph
        hp._p.getparent().remove(hp._p)

        ftr = sec.footer; ftr.is_linked_to_previous = False
        fp = ftr.paragraphs[0]; fp.text = ""
        fl = ftr.add_paragraph(); para_format(fl, line=1.0, after=4); add_top_border(fl, C["rule"], sz=4, space=4)
        set_run_font(fl.add_run(""), F["regular"], 2, C["ink"])
        t = ftr.add_table(rows=1, cols=3, width=self.content_width)
        table_no_borders(t); table_zero_margins(t); table_fixed(t, [Inches(2.0), Inches(CONTENT_IN - 4.0), Inches(2.0)])
        cells = t.rows[0].cells
        for c in cells:
            cell_margins(c, top=0, bottom=0, left=0, right=0)
        p = cells[0].paragraphs[0]; para_format(p, line=1.2)
        r = p.add_run("C9D"); set_run_font(r, F["regular"], 10, C["stamp"])
        r = p.add_run(" Consulting"); set_run_font(r, F["regular"], 10, C["ink"])
        p = cells[1].paragraphs[0]; para_format(p, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(BRAND["attribution"]); set_run_font(r, F["regular"], 9, C["ink_mute"])
        p = cells[2].paragraphs[0]; para_format(p, line=1.2, align=WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run("c9d.consulting"); set_run_font(r, F["mono"], 9, C["ink_mute"], spacing=0.04)
        fp._p.getparent().remove(fp._p)

        # first page: empty header, wordmark + attribution in the footer
        fh = sec.first_page_header; fh.is_linked_to_previous = False; fh.paragraphs[0].text = ""
        ff = sec.first_page_footer; ff.is_linked_to_previous = False
        p = ff.paragraphs[0]; para_format(p, line=1.2, before=0, after=2)
        r = p.add_run("C9D"); set_run_font(r, F["regular"], 16, C["stamp"], spacing=-0.02)
        r = p.add_run(" Consulting"); set_run_font(r, F["regular"], 16, C["ink_bright"], spacing=-0.02)
        p = ff.add_paragraph(); para_format(p, line=1.2)
        r = p.add_run(BRAND["attribution"]); set_run_font(r, F["regular"], 10, C["ink_mute"])
        p = ff.add_paragraph(); para_format(p, line=1.2, before=6)
        sig = "\u2192 " + BRAND["closing"].rstrip(".").upper()
        r = p.add_run(sig); set_run_font(r, F["mono"], 8, C["stamp_dim"], spacing=0.18)
        self.chars.update(sig)
        self.chars.update("C9D Consulting c9d.consulting §·/0123456789" + BRAND["attribution"])
        self.chars.update(self.meta.get("file_id", "") + self.meta.get("classification", ""))

    def build_cover(self):
        m = self.meta
        p = self.doc.add_paragraph(); para_format(p, before=0, line=1.3)
        r = p.add_run(m.get("file_id", "")); set_run_font(r, F["mono"], 11, C["ink_mute"], spacing=0.04)
        p = self.doc.add_paragraph(); para_format(p, line=1.3, after=0)
        r = p.add_run(m.get("classification", "")); set_run_font(r, F["medium"], 9, C["stamp"], caps=True, spacing=0.18)
        p = self.doc.add_paragraph(); para_format(p, before=COVER_GAP_PT, after=10, line=1.05)
        r = p.add_run(m.get("title", "")); set_run_font(r, F["light"], 36, C["ink_bright"], spacing=-0.03)
        p = self.doc.add_paragraph(); para_format(p, line=1.0, after=0, right_indent=Inches(CONTENT_IN - 2.4))
        add_bottom_border(p, C["rule_bright"], sz=8, space=0); set_run_font(p.add_run(""), F["regular"], 2, C["ink"])
        p = self.doc.add_paragraph(); para_format(p, before=16, line=1.3)
        r = p.add_run(m.get("subtitle", "")); set_run_font(r, F["light"], 18, C["ink_mute"])
        # metadata table
        p = self.doc.add_paragraph(); para_format(p, before=COVER_GAP_PT, after=0, line=1.0)
        set_run_font(p.add_run(""), F["regular"], 2, C["ink"])
        rows = [("DOCUMENT", m.get("title", "")), ("PREPARED FOR", m.get("prepared_for", "")),
                ("PREPARED BY", m.get("prepared_by", "")), ("CLASSIFICATION", m.get("classification", "")),
                ("VERSION", m.get("version", ""))]
        t = self.doc.add_table(rows=len(rows), cols=2); table_no_borders(t)
        table_zero_margins(t); table_fixed(t, [Inches(1.5), Inches(CONTENT_IN - 1.5)])
        for i, (k, v) in enumerate(rows):
            row = t.rows[i]; row_no_split(row)
            for j, c in enumerate(row.cells):
                cell_margins(c, top=100, bottom=100, left=0, right=200)
                cell_borders(c, bottom=(C["rule"], 4))
                pp = c.paragraphs[0]; para_format(pp, line=1.3, keep_next=(i < len(rows) - 1))  # v1.4: cover metadata never splits
                if j == 0:
                    rr = pp.add_run(k); set_run_font(rr, F["mono"], 10, C["ink_mute"], spacing=0.10)
                else:
                    rr = pp.add_run(v); set_run_font(rr, F["regular"], 12, C["ink"])
                self.chars.update(k + v)
        for s in (m.get("title", ""), m.get("subtitle", "")):
            self.chars.update(s)
        # page break
        p = self.doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); para_format(p, line=1.0)

    # -- body blocks --
    def heading(self, level, text):
        if level == 1:
            p = self.doc.add_paragraph(); para_format(p, before=0, after=0, line=1.1, keep_next=True, right_indent=BODY_MEASURE_INDENT)
            r = p.add_run(text); set_run_font(r, F["light"], 24, C["ink_bright"], spacing=-0.02)
            p2 = self.doc.add_paragraph(); para_format(p2, after=12, line=1.0, keep_next=True); p2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; p2.paragraph_format.line_spacing = Pt(RULE_GAP_PT)  # v1.3: tight title-to-rule gap
            add_bottom_border(p2, C["rule_bright"], sz=4, space=0); set_run_font(p2.add_run(""), F["regular"], 2, C["ink"])
        elif level == 2:
            eyebrow, title = (text.split(" · ", 1) + [None])[:2] if " · " in text else (None, text)
            p = self.doc.add_paragraph(); para_format(p, before=28, after=4, line=1.3, keep_next=True)
            if eyebrow:
                if re.fullmatch(r"\d+", eyebrow):
                    eyebrow = f"SECTION {int(eyebrow):02d}"
                r = p.add_run(eyebrow); set_run_font(r, F["mono_bold"], 10, C["stamp"], caps=True, spacing=0.18)
            else:
                r = p.add_run("SECTION"); set_run_font(r, F["mono_bold"], 10, C["stamp"], caps=True, spacing=0.18)
            p = self.doc.add_paragraph(); para_format(p, before=0, after=0, line=1.1, keep_next=True, right_indent=BODY_MEASURE_INDENT)
            r = p.add_run(title); set_run_font(r, F["light"], 24, C["ink_bright"], spacing=-0.02)
            p2 = self.doc.add_paragraph(); para_format(p2, after=12, line=1.0, keep_next=True); p2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; p2.paragraph_format.line_spacing = Pt(RULE_GAP_PT)  # v1.3: tight title-to-rule gap
            add_bottom_border(p2, C["rule_bright"], sz=4, space=0); set_run_font(p2.add_run(""), F["regular"], 2, C["ink"])
        elif level == 3:
            p = self.doc.add_paragraph(); para_format(p, before=18, after=6, line=1.2, keep_next=True, right_indent=BODY_MEASURE_INDENT)
            r = p.add_run(text); set_run_font(r, F["light"], 18, C["ink_bright"], spacing=-0.01)
        else:
            p = self.doc.add_paragraph(); para_format(p, before=12, after=4, line=1.3, keep_next=True, right_indent=BODY_MEASURE_INDENT)
            r = p.add_run(text); set_run_font(r, F["medium"], 14, C["ink_bright"])
        self.chars.update(text)

    def paragraph(self, text):
        p = self.doc.add_paragraph(); para_format(p, after=10, line=1.6, right_indent=BODY_MEASURE_INDENT, keep_together=False)
        if text.startswith("*") and text.endswith("*") and not text.startswith("**"):
            # standalone description line under the title: muted body, not an italic block
            r = p.add_run(text[1:-1]); set_run_font(r, F["light"], 12, C["ink_mute"]); self.chars.update(text)
            return
        if len(text) <= 60 or getattr(self, "next_is_table", False):
            p.paragraph_format.keep_with_next = True
        self.runs(p, text)

    def list_block(self, items, ordered):
        for n, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            para_format(p, after=6, line=1.5, right_indent=BODY_MEASURE_INDENT, left_indent=Inches(0.4), hanging=Inches(0.4))
            marker = f"{n:02d}." if ordered else "–"
            r = p.add_run(marker + "\t"); set_run_font(r, F["mono"] if ordered else F["regular"], 10 if ordered else BODY_PT, C["stamp_dim"])
            self.chars.update(marker)
            self.runs(p, it)
            # tab stop at the hanging indent
            pPr = p._p.get_or_add_pPr(); tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "left"); tab.set(qn("w:pos"), str(int(0.4 * 1440))); tabs.append(tab); pPr.append(tabs)

    def table(self, rows):
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        # proportion columns to content, header sets the floor
        weights = []
        for j in range(ncols):
            body_max = max([len(re.sub(r"[*`\[\]]", "", r[j])) for r in rows[1:]] or [0])
            hdr = len(rows[0][j])
            w = max(hdr * 1.4 + 3, min(body_max, 60) ** 0.85, 6)
            weights.append(w)
        tot = sum(weights)
        widths = [Inches(CONTENT_IN * w / tot) for w in weights]
        t = self.doc.add_table(rows=len(rows), cols=ncols); table_no_borders(t); table_fixed(t, widths)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), str(CELL_PAD_H)); ind.set(qn("w:type"), "dxa"); t._tbl.tblPr.append(ind)
        tcm = OxmlElement("w:tblCellMar")
        for side, v in (("top", CELL_PAD_TOP), ("left", CELL_PAD_H), ("bottom", CELL_PAD_BOTTOM), ("right", CELL_PAD_H)):
            e = OxmlElement("w:" + side); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); tcm.append(e)
        t._tbl.tblPr.append(tcm)
        for i, row in enumerate(rows):
            tr = t.rows[i]; row_no_split(tr, header=(i == 0))
            for j, val in enumerate(row):
                c = tr.cells[j]; last = (j == ncols - 1)
                cell_margins(c, top=CELL_PAD_TOP, bottom=CELL_PAD_BOTTOM, left=CELL_PAD_H, right=CELL_PAD_H)
                pp = c.paragraphs[0]; para_format(pp, line=CELL_LINE, keep_next=(len(rows) <= 8 and i < len(rows) - 1))
                cell_valign(c, "center" if i == 0 else "top")
                if i == 0:
                    cell_shade(c, C["bg2"]); cell_borders(c, bottom=(C["rule_bright"], 8))
                    rr = pp.add_run(val); set_run_font(rr, F["mono_bold"], 9, C["stamp_dim"], caps=True, spacing=0.18)
                    self.chars.update(val)
                else:
                    cell_borders(c, bottom=(C["rule"], 4))
                    for kind, tx in inline_runs(val):
                        self.run(pp, tx, kind, 10)
        sp = self.doc.add_paragraph(); para_format(sp, after=8, line=1.0); set_run_font(sp.add_run(""), F["regular"], 4, C["ink"])

    def build(self, blocks):
        self.build_header_footer()
        self.build_cover()
        for idx, (kind, val) in enumerate(blocks):
            self.next_is_table = idx + 1 < len(blocks) and blocks[idx + 1][0] == "table"
            if kind == "h1":
                self.heading(1, val)
            elif kind == "h2":
                self.heading(2, val)
            elif kind == "h3":
                self.heading(3, val)
            elif kind == "h4":
                self.heading(4, val)
            elif kind == "p":
                self.paragraph(val)
            elif kind == "ul":
                self.list_block(val, False)
            elif kind == "ol":
                self.list_block(val, True)
            elif kind == "table":
                self.table(val)
        self.closing_signature()

    def closing_signature(self):
        """Locked closing line as the closing signature: hairline, then mono caps in stamp (v1.3)."""
        body = self.doc.element.body
        paras = body.findall(qn("w:p"))
        if paras:   # keep the last body paragraph with the signature so it never stands alone
            from docx.text.paragraph import Paragraph
            Paragraph(paras[-1], self.doc).paragraph_format.keep_with_next = True
        p = self.doc.add_paragraph(); para_format(p, before=36, after=0, line=1.0, keep_next=True)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; p.paragraph_format.line_spacing = Pt(RULE_GAP_PT)
        add_bottom_border(p, C["rule"], sz=4, space=0); set_run_font(p.add_run(""), F["regular"], 2, C["ink"])
        p = self.doc.add_paragraph(); para_format(p, before=10, after=0, line=1.2)
        sig = "\u2192 " + BRAND["closing"].rstrip(".").upper()
        r = p.add_run(sig); set_run_font(r, F["mono_bold"], 9, C["stamp"], spacing=0.18)
        self.chars.update(sig)

    def save(self, path, embed=True):
        # settings: embed fonts, update fields on open
        settings = self.doc.settings.element
        if embed:
            e = OxmlElement("w:embedTrueTypeFonts"); settings.insert(0, e)
            e2 = OxmlElement("w:saveSubsetFonts"); settings.insert(1, e2)
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)
        # core properties
        cp = self.doc.core_properties
        cp.title = self.meta.get("title", ""); cp.author = "Brandon Wilburn, C9D Consulting"
        cp.subject = self.meta.get("subtitle", ""); cp.comments = ""
        schema_order(self.doc.element.body)
        for sec in self.doc.sections:
            for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
                schema_order(part._element)
        self.doc.save(path)
        if embed:
            embed_fonts(path, self.chars)

# ---- font embedding ----------------------------------------------------------
def obfuscate(data, guid):
    key = bytes.fromhex(guid.replace("-", "").replace("{", "").replace("}", ""))
    d = bytearray(data)
    for i in range(32):
        d[i] ^= key[15 - (i % 16)]
    return bytes(d)

def subset_font(path, chars):
    f = TTFont(path)
    opts = ftsubset.Options(); opts.name_IDs = ["*"]; opts.notdef_outline = True; opts.layout_features = ["kern"]
    opts.hinting = False; opts.desubroutinize = True
    s = ftsubset.Subsetter(opts)
    text = "".join(sorted(chars)) + "".join(chr(c) for c in range(32, 127))
    s.populate(text=text)
    s.subset(f)
    buf = io.BytesIO(); f.save(buf); return buf.getvalue()

def embed_fonts(path, chars):
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        fonttable = zin.read("word/fontTable.xml").decode("utf8")
        rels_name = "word/_rels/fontTable.xml.rels"
        rels = zin.read(rels_name).decode("utf8") if rels_name in names else \
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>'
        ctypes = zin.read("[Content_Types].xml").decode("utf8")
        font_entries, rel_entries, files = [], [], []
        used_xml = "".join(zin.read(n).decode("utf8") for n in names if n.startswith("word/") and n.endswith(".xml"))
        for n, (family, fname) in enumerate(FONT_FILES.items(), 1):
            if f'w:ascii="{family}"' not in used_xml:
                continue
            data = subset_font(os.path.join(FONT_DIR, fname), chars)
            guid = "{" + str(uuid.uuid4()).upper() + "}"
            rid = f"rIdFont{n}"; target = f"fonts/font{n}.odttf"
            files.append((f"word/{target}", obfuscate(data, guid)))
            italic = "Italic" in fname
            tag = "embedItalic" if italic else "embedRegular"
            font_entries.append(
                f'<w:font w:name="{family}"><w:charset w:val="00"/><w:family w:val="auto"/><w:pitch w:val="variable"/>'
                f'<w:{tag} r:id="{rid}" w:fontKey="{guid}"/></w:font>')
            rel_entries.append(
                f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="{target}"/>')
        # strip any existing declarations of the same families
        for family in FONT_FILES:
            fonttable = re.sub(rf'<w:font w:name="{re.escape(family)}">.*?</w:font>', "", fonttable, flags=re.S)
        if 'xmlns:r=' not in fonttable:
            fonttable = fonttable.replace("<w:fonts ", '<w:fonts xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
        fonttable = fonttable.replace("</w:fonts>", "".join(font_entries) + "</w:fonts>")
        rels = rels.replace("</Relationships>", "".join(rel_entries) + "</Relationships>")
        if 'Extension="odttf"' not in ctypes:
            ctypes = ctypes.replace("<Default ", '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/><Default ', 1)
        for item in zin.infolist():
            if item.filename in ("word/fontTable.xml", rels_name, "[Content_Types].xml"):
                continue
            if item.filename == "word/stylesWithEffects.xml":
                continue
            if item.filename == "word/_rels/document.xml.rels":
                data = zin.read(item.filename).decode("utf8")
                data = re.sub(r'<Relationship [^>]*stylesWithEffects[^>]*/>', "", data)
                zout.writestr(item, data); continue
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/fontTable.xml", fonttable)
        zout.writestr(rels_name, rels)
        zout.writestr("[Content_Types].xml", ctypes)
        for fn, data in files:
            zout.writestr(fn, data)
    shutil.move(tmp, path)

# ---- main --------------------------------------------------------------------
# ---- check ---------------------------------------------------------------------------------
# Voice lists match c9d_html_build.py; keep them in step.
REFUSED = [r"\btransform(ation|ative|s|ed|ing)?\b", r"\bjourney\b", r"\binnovat(ive|ion)\b",
           r"\bempower(ment|s|ed|ing)?\b", r"\bAI maturity model\b", r"\bAI cent(er|re) of excellence\b"]

def docx_text(path):
    """Every text run in the body, headers and footers, in document order."""
    with zipfile.ZipFile(path) as z:
        parts = [n for n in z.namelist() if re.match(r"word/(document|header\d*|footer\d*)\.xml$", n)]
        xml = " ".join(z.read(n).decode("utf8") for n in parts)
        fonts = z.read("word/fontTable.xml").decode("utf8")
    xml = re.sub(r"<w:instrText[^>]*>.*?</w:instrText>", " ", xml, flags=re.S)
    xml = re.sub(r"</w:p>", "\n", xml)
    import html as _html
    return _html.unescape(re.sub(r"<[^>]+>", "", xml)), fonts

def xml_all(path):
    with zipfile.ZipFile(path) as z:
        return " ".join(z.read(n).decode("utf8") for n in z.namelist()
                        if re.match(r"word/(document|header\d*|footer\d*|styles)\.xml$", n))

def md_words(body):
    t = re.sub(r"(?m)^#\s+.*$", " ", body)                 # title lives on the cover
    t = re.sub(r"(?m)^\s*\|?\s*:?-{3,}.*$", " ", t)        # table separators
    t = re.sub(r"\[([^\]]+)\]\((?:[^)]+)\)", r"\1", t)     # link targets
    t = re.sub(r"(?m)^\s*\d+\.\s+", " ", t)                # list numbers render as 01.
    t = re.sub(r"(?m)^(##\s+)\d+\s*\u00b7\s+", r"\1", t)   # numeric labels become SECTION 0N
    import collections
    return collections.Counter(w.lower() for w in re.findall(r"[A-Za-z0-9$%']+", t))

def check(src, out, draft=False):
    import collections
    meta, body = parse_front_matter(open(src, encoding="utf8").read())
    text, fonts = docx_text(out)
    problems = []
    have = collections.Counter(w.lower() for w in re.findall(r"[A-Za-z0-9$%']+", text))
    miss = md_words(body) - have
    if miss:
        problems.append("parity: source words missing from the .docx: " +
                        ", ".join(f"{w}x{n}" for w, n in miss.most_common(20)))
    embedded = set(re.findall(r'<w:font w:name="([^"]+)"><w:charset[^>]*/><w:family[^>]*/><w:pitch[^>]*/><w:embed', fonts))
    used = set(re.findall(r'w:ascii="([^"]+)"', xml_all(out))) & set(F.values())
    if used - embedded:
        problems.append(f"fonts: not embedded: {sorted(used - embedded)} (built with --no-embed?)")
    attr = BRAND["attribution"]
    n = text.count(attr)
    stray = re.findall(re.escape(attr) + r"(?=[.,;:])", text)
    if n < 2: problems.append(f"chrome: attribution '{attr}' on cover and footer expected, found {n}")
    if stray: problems.append(f"chrome: attribution carries trailing punctuation x{len(stray)}")
    sig = BRAND["closing"].rstrip(".").upper()
    if text.count(sig) < 2: problems.append("chrome: closing signature missing from cover or last page")
    prose = body + "\n" + meta.get("title", "") + "\n" + meta.get("subtitle", "")
    if "\u2014" in prose: problems.append(f"voice: {prose.count(chr(0x2014))} em-dash(es) in source")
    bare = re.findall(r"\bC9D\b(?! (?i:consulting))(?!-)", prose)   # legal all-caps "C9D CONSULTING LLC" is the full form
    if bare: problems.append(f"voice: bare 'C9D' used {len(bare)} time(s); write 'C9D Consulting'")
    for pat in REFUSED:
        if re.search(pat, prose, re.I): problems.append(f"voice: refused vocabulary /{pat}/")
    brackets = re.findall(r"(?<!\])\[[^\]]+\](?!\()", body)
    if brackets and not draft:
        problems.append(f"placeholders: {len(brackets)} unresolved, e.g. {brackets[0]} (use --draft for templates)")
    for p in problems: print("FAIL ", p)
    if not problems: print("PASS  parity, fonts, chrome, voice, placeholders")
    return 1 if problems else 0

# ---- main --------------------------------------------------------------------------------
def main():
    import argparse
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source"); ap.add_argument("out")
    ap.add_argument("--no-embed", action="store_true", help="skip font embedding (quick drafts and the substitution stress test only)")
    ap.add_argument("--tokens", help="path to tokens.json (default: the brand repo copy, else fetched)")
    ap.add_argument("--refresh-tokens", action="store_true", help="re-fetch tokens.json from the brand repo")
    ap.add_argument("--check", action="store_true", help="verify parity, fonts, chrome and voice after building")
    ap.add_argument("--draft", action="store_true", help="with --check, allow unresolved [placeholders]")
    a = ap.parse_args()
    load_tokens(a.tokens, a.refresh_tokens)
    print(f"tokens: {BRAND['source']} (v{BRAND['token_version']})")
    embed = not a.no_embed
    text = open(a.source, encoding="utf8").read()
    meta, body = parse_front_matter(text)
    blocks = parse_blocks(body)
    os.makedirs(FONT_DIR, exist_ok=True)
    if embed:
        ensure_fonts()
    w = Writer(meta); w.build(blocks); w.save(a.out, embed=embed)
    print(f"wrote {a.out} (layout v{VERSION}) {os.path.getsize(a.out) // 1024} KB")
    if a.check:
        sys.exit(check(a.source, a.out, draft=a.draft))

if __name__ == "__main__":
    main()
